from __future__ import annotations

import ast
import base64
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE_ROOT = Path(
    r"C:\Users\张可姝\Desktop\FIT5196 Data Wrangling\Ass1"
    r"\Group050_A1-20260831T145411Z-1-001\Group050_A1"
)
SOLUTION_PATH = SOURCE_ROOT / "Group050_solution.ipynb"
EDA_PATH = SOURCE_ROOT / "Group050_EDA.ipynb"
TEXT_FUNCTIONS_PATH = SOURCE_ROOT / "Group050_A1_submission" / "Group050_text_functions.py"
OUTPUT_DIR = Path(r"C:\monash\Monash_Project\5196_Project1\outputs")
OUTPUT_PATH = OUTPUT_DIR / "Group050_Keshu_Contribution_Bilingual_Walkthrough_Task2_to_6.docx"
ASSET_DIR = Path(r"C:\monash\Monash_Project\5196_Project1\.codex_tmp\keshu_walkthrough_assets")


NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GREY = "F2F2F2"
MID_GREY = "D9D9D9"
TEXT_GREY = "404040"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = MID_GREY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", east_asia="Microsoft YaHei", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text: str, style=None, bold_lead: str | None = None, keep_with_next=False):
    paragraph = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.keep_with_next = keep_with_next
    return paragraph


def add_bilingual(doc, zh: str, en: str) -> None:
    p_zh = doc.add_paragraph()
    r_zh = p_zh.add_run(zh)
    set_run_font(r_zh, east_asia="Microsoft YaHei", size=10.5)
    p_zh.paragraph_format.space_after = Pt(1.5)
    p_zh.paragraph_format.line_spacing = 1.12
    p_en = doc.add_paragraph()
    r_en = p_en.add_run(en)
    set_run_font(r_en, size=9.8, color=TEXT_GREY)
    p_en.paragraph_format.space_after = Pt(7)
    p_en.paragraph_format.line_spacing = 1.08


def add_bullet(doc, zh: str, en: str | None = None, level=0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    r = p.add_run(zh)
    set_run_font(r, size=10.2)
    if en:
        r2 = p.add_run(f" / {en}")
        set_run_font(r2, size=9.7, color=TEXT_GREY)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08


def add_heading(doc, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        set_run_font(
            run,
            size={1: 17, 2: 13.5, 3: 11.5}.get(level, 10.5),
            bold=True,
            color="000000",
        )


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = str(header)
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, size=9.2, bold=True, color=WHITE)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            set_cell_shading(cells[j], PALE_BLUE if i % 2 else WHITE)
            set_cell_borders(cells[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.03
                if j == 0 and len(headers) <= 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=font_size)
        if widths:
            for j, width in enumerate(widths):
                cells[j].width = Inches(width)
    for row in table.rows:
        if widths:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, code: str) -> None:
    for line in code.rstrip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_GREY)
        p_pr.append(shd)
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", east_asia="Microsoft YaHei", size=8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def extract_function_code(path: Path, names: list[str]) -> dict[str, str]:
    source = path.read_text(encoding="utf-8")
    tree = ast.parse(source)
    lines = source.splitlines()
    result = {}
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in names:
            result[node.name] = "\n".join(lines[node.lineno - 1 : node.end_lineno])
    return result


def extract_eda_figures() -> tuple[Path, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    notebook = json.loads(EDA_PATH.read_text(encoding="utf-8"))
    pngs = []
    for output in notebook["cells"][20].get("outputs", []):
        data = output.get("data", {})
        if "image/png" in data:
            payload = data["image/png"]
            if isinstance(payload, list):
                payload = "".join(payload)
            pngs.append(base64.b64decode(payload))
    if len(pngs) != 2:
        raise RuntimeError(f"Expected two Figure 7/8 PNGs, found {len(pngs)}")
    paths = (ASSET_DIR / "figure7.png", ASSET_DIR / "figure8.png")
    for path, raw in zip(paths, pngs):
        path.write_bytes(raw)
    return paths


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=TEXT_GREY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Normal"

    if "Code" not in [s.name for s in doc.styles]:
        code_style = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(8.2)

    for sec in doc.sections:
        add_page_number(sec.footer.paragraphs[0])


def build() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    figure7_path, figure8_path = extract_eda_figures()
    function_code = extract_function_code(
        TEXT_FUNCTIONS_PATH,
        ["build_latin_analysis", "contains_non_latin_script"],
    )

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(110)
    title.paragraph_format.space_after = Pt(16)
    r = title.add_run("Group050 Keshu Contribution Walkthrough")
    set_run_font(r, name="Aptos Display", size=25, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Task 2 to Task 6  中英双语代码与分析说明")
    set_run_font(sr, size=14, bold=True, color=NAVY)
    subtitle.paragraph_format.space_after = Pt(28)

    add_table(
        doc,
        ["Item", "Details"],
        [
            ["Member", "Keshu Zhang  36436763"],
            ["Assigned domain", "Products and product reviews"],
            ["Task 3 ownership", "build_latin_analysis and contains_non_latin_script"],
            ["EDA ownership", "Figures 7–8  Findings 9–10  MLQ 4–5"],
            ["Purpose", "Individual understanding and group handover  个人复习与小组交接"],
        ],
        widths=[1.7, 5.2],
        font_size=10,
    )

    add_bilingual(
        doc,
        "本说明按照小组任务划分整理本人负责的代码、验证、EDA图表、发现与机器学习问题。它用于解释代码和证明个人理解，不替代正式提交文件。",
        "This walkthrough explains the code, validation, EDA figures, findings and machine-learning questions assigned to me. It supports individual understanding and team handover; it is not a substitute for the submitted notebooks and CSV files.",
    )

    doc.add_page_break()
    add_heading(doc, "1 Personal Scope and End to End Flow  个人范围与整体流程", 1)
    add_table(
        doc,
        ["Task", "My responsibility  我的职责", "Current evidence  当前证据"],
        [
            ["Task 2", "Build products and product_reviews", "1,000 × 21 and 7,000 × 21"],
            ["Task 3.4", "Multilingual text functions", "2 public functions integrated"],
            ["Task 4", "Products reviews text and final validation integration", "20 checks PASS"],
            ["Task 5", "Figures 7 and 8", "Both reproduced in EDA notebook"],
            ["Task 6", "Findings 9–10 and MLQ 4–5", "Aligned to Figures 7–8"],
        ],
        widths=[0.8, 3.7, 2.5],
    )
    add_bilingual(
        doc,
        "我的工作链条是：先依赖Task 1得到结构化XML和JSON记录及数据字典；再建立两张标准化表；随后执行字段、主外键、重复、文本和多语言验证；最后用两张表完成Figures 7–8，并据此写Findings 9–10和MLQ 4–5。",
        "My workflow begins with the structured JSON/XML records and public dictionary created in Task 1. I then construct the two standardised tables, validate schema, keys, duplication, text and multilingual fields, and finally use the tables for Figures 7–8, Findings 9–10 and MLQ 4–5.",
    )
    add_heading(doc, "1.1 Dependencies  对其他成员工作的依赖", 2)
    add_table(
        doc,
        ["Dependency", "Owner", "How my component uses it"],
        [
            ["clean_narrative_text", "Orders and order items owner", "Cleans product descriptions and review bodies"],
            ["extract_order_reference", "Customers owner", "Extracts an embedded HORD or CORD reference before cleaning"],
            ["extract_product_sku", "Deliveries owner", "Extracts an embedded SKU before cleaning"],
            ["orders order_items customers", "Other Task 2 owners", "Required for product-review foreign-key validation"],
            ["All component validation rows", "All members", "Required before final validation-register integration"],
        ],
        widths=[1.7, 2.0, 3.3],
        font_size=8.9,
    )
    add_bilingual(
        doc,
        "重要边界：调用共享函数不等于拥有该函数。我的个人Task 3贡献仅为build_latin_analysis和contains_non_latin_script。",
        "Important ownership boundary: calling a shared function does not mean that I authored it. My personal Task 3 contribution is limited to build_latin_analysis and contains_non_latin_script.",
    )

    doc.add_page_break()
    add_heading(doc, "2 Task 2 Products Table  产品表", 1)
    add_bilingual(
        doc,
        "目标粒度是一行对应一个product_id。产品数据来自结构化解析后的XML ProductCatalogue/Product记录，输出字段顺序由public_data_dictionary.csv动态取得。",
        "The target grain is one row per product_id. Product data comes from the structurally parsed XML ProductCatalogue/Product records, and the output field order is read dynamically from public_data_dictionary.csv.",
    )
    add_heading(doc, "2.1 Processing sequence  处理顺序", 2)
    for zh, en in [
        ("检查Task 1生成的XML产品记录和数据字典对象是否存在。", "Check that the Task 1 XML product records and dictionary objects exist."),
        ("读取数据字典中的21个产品目标字段及顺序。", "Read the 21 Product target columns and their order from the public dictionary."),
        ("逐条标准化文本、货币、整数、数值、日期和布尔值。", "Standardise text, currency, integer, numeric, date and boolean values record by record."),
        ("清洗product_description_clean，并优先调用已加载的共享clean_narrative_text。", "Create product_description_clean, delegating to the shared clean_narrative_text when available."),
        ("按product_id分组，比较重复记录的各字段，而不是静默选择某个来源。", "Group by product_id and compare duplicate records field by field rather than silently choosing one value."),
        ("若发现多个不同的非空标准化值，则记录冲突并停止。", "Record and stop on multiple distinct non-missing normalised values."),
        ("保留精确字段顺序，按product_id稳定排序并执行即时构造检查。", "Preserve exact column order, sort deterministically by product_id and run immediate construction guards."),
    ]:
        add_bullet(doc, zh, en)

    add_heading(doc, "2.2 Main helper functions  主要辅助函数", 2)
    add_table(
        doc,
        ["Function", "What it does  作用", "Why it is needed  原因"],
        [
            ["_normalise_required_product_text", "Trim and NFC-normalise required text; reject missing/empty input", "Preserves leading zeros and category/identifier case"],
            ["_parse_product_currency", "Remove AUD and commas; Decimal conversion; 2-decimal rounding", "Prevents currency labels from corrupting numeric conversion"],
            ["_parse_product_integer", "Convert only mathematically integral values", "Avoids silently rounding invalid values"],
            ["_parse_product_number", "Convert measurements through Decimal to float", "Retains published measurement precision"],
            ["_parse_product_date", "DD/MM/YYYY to YYYY-MM-DD", "Matches target date format"],
            ["_parse_product_boolean", "Y/N to True/False", "Matches target boolean contract"],
            ["_clean_product_description_for_task2", "Shared cleaner or equivalent private fallback", "Allows Task 2 to build while keeping Task 3 ownership separate"],
            ["_standardise_product_record", "Map one XML Product to exactly 21 target fields", "Maintains one-record one-product grain"],
            ["_reconcile_product_candidates", "Compare duplicates by product_id and field", "Exposes conflicts instead of applying source precedence"],
        ],
        widths=[2.05, 2.55, 2.4],
        font_size=8.35,
    )
    add_heading(doc, "2.3 Output and issues handled  输出与问题处理", 2)
    add_table(
        doc,
        ["Evidence", "Observed result", "Interpretation"],
        [
            ["Rows and columns", "1,000 rows × 21 columns", "One canonical row per product"],
            ["Primary key", "0 missing and 0 duplicated product_id", "Product grain is preserved"],
            ["Reconciliation", "0 field-level conflicts", "Normalised duplicates agree"],
            ["Missing strings", "0 blank cells; 0 unintended missing values", "Current source happens not to require a literal NaN in product strings"],
            ["Schema", "Exact data-dictionary columns and order", "No helper or EDA-only columns exported"],
        ],
        widths=[1.5, 2.2, 3.3],
        font_size=8.8,
    )

    add_heading(doc, "3 Task 2 Product Reviews Table  产品评论表", 1)
    add_bilingual(
        doc,
        "目标粒度是一行对应一个canonical review_id。评论同时来自JSON和XML，因此先分别标准化，再合并候选记录，并按review_id进行字段级协调。",
        "The target grain is one row per canonical review_id. Reviews occur in both JSON and XML, so each source is standardised separately, the candidates are combined, and rows are reconciled field by field using review_id.",
    )
    add_heading(doc, "3.1 Correct order of text operations  文本操作的正确顺序", 2)
    add_table(
        doc,
        ["Step", "Operation", "Reason"],
        [
            ["1", "Read raw review text", "References still exist in the raw string"],
            ["2", "extract_order_reference and extract_product_sku", "Extraction must happen before the reference wrapper is removed"],
            ["3", "clean_narrative_text", "Produce review_body_clean while preserving multilingual text"],
            ["4", "build_latin_analysis(review_body_clean)", "Create the separate Latin-only analytical representation"],
            ["5", "contains_non_latin_script(review_body_clean)", "Create the non-Latin-script indicator"],
            ["6", "Derive character and whitespace-token counts", "Both measures must be based on review_body_clean"],
        ],
        widths=[0.55, 3.2, 3.45],
        font_size=8.9,
    )
    add_bilingual(
        doc,
        "顺序不能交换。若先执行clean_narrative_text，完整的Reference和SKU wrapper会被删除，之后将无法提取embedded reference。",
        "The order is not interchangeable. If clean_narrative_text runs first, the complete Reference and SKU wrapper is removed and the embedded references can no longer be extracted.",
    )
    add_heading(doc, "3.2 Source differences and standardisation  来源差异", 2)
    add_table(
        doc,
        ["Field type", "JSON representation", "XML representation", "Target"],
        [
            ["Field names", "camelCase", "Title_Case", "snake_case"],
            ["Timestamp", "YYYY-MM-DD HH:MM:SS", "DD/MM/YYYY HH:MM:SS", "YYYY-MM-DD HH:MM:SS"],
            ["Boolean", "true or false", "Y or N", "True or False"],
            ["Integer", "JSON number", "XML text", "Python integer then CSV numeric text"],
            ["Narrative", "Raw multilingual reviewText", "Raw multilingual Review_Text", "Clean text plus derived fields"],
        ],
        widths=[1.25, 1.85, 1.85, 2.25],
        font_size=8.7,
    )
    add_heading(doc, "3.3 Reconciliation and row flow  重复与重叠处理", 2)
    add_table(
        doc,
        ["Stage", "Rows or keys", "Meaning"],
        [
            ["Structured JSON rows", "3,946", "Includes 96 within-source duplicate rows"],
            ["Structured XML rows", "3,946", "Includes 96 within-source duplicate rows"],
            ["Unique keys per source", "3,850 each", "Duplicates removed conceptually by review_id"],
            ["Cross-source overlap", "700 review_id values", "Same canonical review appears in both sources"],
            ["Combined candidates", "7,892", "All normalised rows before reconciliation"],
            ["Union of review keys", "7,000", "Expected canonical key set"],
            ["Final output", "7,000 rows × 21 columns", "One canonical row per review_id"],
            ["Field-level conflicts", "0", "Overlapping normalised rows agree"],
        ],
        widths=[2.1, 1.45, 3.55],
        font_size=8.75,
    )

    add_heading(doc, "4 Task 3 4 Multilingual Functions  多语言函数", 1)
    add_bilingual(
        doc,
        "这两个函数是我的明确Task 3个人贡献。它们是纯函数：不读写文件、不访问网络、不依赖特定行号，只根据输入返回确定结果。",
        "These two functions are my explicit individual Task 3 contribution. They are pure functions: they perform no file I/O or network access, use no row-specific lookup, and return a deterministic result for the supplied input.",
    )

    add_heading(doc, "4.1 build_latin_analysis", 2)
    add_code(doc, function_code["build_latin_analysis"])
    add_table(
        doc,
        ["Code stage", "中文理解", "English explanation"],
        [
            ["Input guard", "非字符串或字面量NaN直接返回NaN", "Reject unsupported input and preserve the literal sentinel"],
            ["NFC normalisation", "统一组合字符表示，例如带变音符号的Latin字母", "Normalise canonically equivalent Unicode sequences"],
            ["Unicode category L", "只把字符类别为Letter的字符当作字母", "Treat Unicode Letter categories as letters"],
            ["LATIN in Unicode name", "保留Latin字母，包括é等字符", "Retain Latin-script letters including diacritics"],
            ["Other-script letter", "删除非Latin字母并插入空格，防止前后Latin词被错误拼接", "Remove non-Latin letters and insert a separator to avoid joining surrounding Latin words"],
            ["Combining marks", "只有紧跟已保留Latin字母时才保留", "Keep a combining mark only when it belongs to retained Latin text"],
            ["Digits punctuation spaces", "保留数字、标点和其他非字母字符", "Retain permitted non-letter characters"],
            ["Final result", "合并空格；没有Latin字母则返回NaN", "Collapse whitespace; return NaN when no Latin letter remains"],
        ],
        widths=[1.4, 2.8, 2.8],
        font_size=8.35,
    )
    add_bilingual(
        doc,
        "关键原理是使用unicodedata.category和unicodedata.name判断脚本，而不是使用ASCII范围。这样é会被正确保留，中文或日文字母会被删除。",
        "The key principle is script-aware Unicode classification through unicodedata.category and unicodedata.name rather than an ASCII range test. This preserves characters such as é while removing Chinese or Japanese letters.",
    )

    doc.add_page_break()
    add_heading(doc, "4.2 contains_non_latin_script", 2)
    add_code(doc, function_code["contains_non_latin_script"])
    add_table(
        doc,
        ["Code stage", "中文理解", "English explanation"],
        [
            ["Input guard", "非字符串或NaN返回False", "Unsupported or sentinel input contains no usable non-Latin script"],
            ["NFC normalisation", "保证Unicode表示一致", "Use a consistent Unicode representation"],
            ["category starts with L", "只检查字母，不把emoji和符号误判为脚本", "Inspect letters only so emoji and symbols are not treated as a script"],
            ["LATIN not in name", "只要发现一个非Latin字母就返回True", "Return True when at least one non-Latin letter is found"],
            ["any generator", "找到第一个满足条件的字符即可停止", "Short-circuit after the first matching character"],
        ],
        widths=[1.45, 2.8, 2.75],
        font_size=8.6,
    )
    add_heading(doc, "4.3 Boundary examples  边界示例", 2)
    add_table(
        doc,
        ["Input", "build_latin_analysis", "contains_non_latin_script", "Reason"],
        [
            ["café 2024!", "café 2024!", "False", "é is Latin script"],
            ["产品 good", "good", "True", "Chinese letters are removed from analysis but detected"],
            ["中文", "NaN", "True", "No Latin letter remains"],
            ["😊 123", "NaN", "False", "Emoji and digits are not non-Latin letters"],
            ["NaN", "NaN", "False", "Literal missing-string sentinel is preserved"],
            ["None", "NaN", "False", "Unsupported missing input follows the interface"],
        ],
        widths=[1.15, 2.05, 1.7, 2.1],
        font_size=8.7,
    )

    doc.add_page_break()
    add_heading(doc, "5 Task 4 Validation  验证", 1)
    add_bilingual(
        doc,
        "我的范围共有20条正式验证记录，全部为PASS。每条记录包含稳定ID、检查说明、实际观测值、PASS或FAIL、evidence以及resolution or interpretation。",
        "My scope contains 20 formal validation rows and all currently pass. Every row includes a stable ID, check description, observed result, PASS/FAIL status, evidence and a resolution or interpretation.",
    )
    add_bilingual(
        doc,
        "PASS不是硬编码结果。状态由当前数据计算产生；如果检查失败，记录应保留为FAIL并给出调查或修复方向。",
        "PASS is not hard-coded. Status is derived from the current data; if a check fails, the row must remain visible as FAIL with an investigation or resolution path.",
    )
    validation_rows = [
        ["VAL-PROD-SCHEMA-01", "Exact Product columns and order", "PASS"],
        ["VAL-PROD-TYPE-01", "Product semantic types and required missingness", "PASS"],
        ["VAL-PROD-PK-01", "product_id complete and unique; 1,000 rows", "PASS"],
        ["VAL-PROD-RANGE-01", "Price cost weight warranty date year and SKU invariants", "PASS"],
        ["VAL-PROD-CAT-01", "Product categories supported by parsed XML values", "PASS"],
        ["VAL-PROD-FLOW-01", "1,000 XML rows to 1,000 canonical keys", "PASS"],
        ["VAL-PROD-RECON-01", "0 unresolved Product field conflicts", "PASS"],
        ["VAL-REV-SCHEMA-01", "Exact Product Review columns and order", "PASS"],
        ["VAL-REV-TYPE-01", "Review semantic types and required missingness", "PASS"],
        ["VAL-REV-PK-01", "review_id complete and unique; 7,000 rows", "PASS"],
        ["VAL-REV-FK-01", "Order item product and customer foreign keys resolve", "PASS"],
        ["VAL-REV-RANGE-01", "Rating votes and derived counts in valid ranges", "PASS"],
        ["VAL-REV-CAT-01", "Review categories and language codes supported by sources", "PASS"],
        ["VAL-REV-FLOW-01", "7,892 candidates reconcile to 7,000 union keys", "PASS"],
        ["VAL-REV-RECON-01", "0 unresolved Review field conflicts", "PASS"],
        ["VAL-TEXT-TEST-01", "18/18 public 21/21 student and 6/6 interface tests", "PASS"],
        ["VAL-TEXT-DERIVED-01", "0 length and word-count mismatches", "PASS"],
        ["VAL-TEXT-REF-01", "0 invalid or inconsistent order and SKU references", "PASS"],
        ["VAL-TEXT-MULTI-01", "287 reviews flagged non-Latin; 0 derivative mismatches", "PASS"],
        ["VAL-TEXT-NAN-01", "0 blank or Python-missing required strings", "PASS"],
    ]
    add_table(doc, ["Validation ID", "Observed check", "Status"], validation_rows, widths=[1.65, 4.75, 0.6], font_size=8.0)

    add_heading(doc, "6 Task 5 Figures 7 and 8  图表", 1)
    add_heading(doc, "6.1 Figure 7 Review rating and text behaviour", 2)
    add_bilingual(
        doc,
        "分析问题：评论长度和helpful votes如何随评分与语言组变化？观察单位为一条canonical product review；分母为全部7,000条评论，其中英语6,311条、非英语689条。只使用product_reviews，不需要join。",
        "Analytical question: How do review length and helpful votes vary across rating and language groups? The unit is one canonical product review. The denominator is all 7,000 reviews: 6,311 English and 689 non-English. The analysis uses product_reviews and requires no join.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(figure7_path), width=Inches(6.75))
    add_bilingual(
        doc,
        "结果：英语评论长度中位数从1星的974字符降至5星约904.5字符；非英语组没有单调趋势。各评分和语言组的helpful votes中位数约为42至49.5。限制：英语占多数，非英语合并了多种脚本；字符长度跨脚本不可完全比较，helpful votes还受曝光时间影响。",
        "Result: median English review length falls from 974 characters at one star to about 904.5 at five stars, while the non-English pattern is not monotonic. Median helpful votes remain about 42–49.5. Limitation: English dominates, non-English combines several scripts, character counts are not fully comparable across scripts, and helpful votes depend on exposure time.",
    )
    add_heading(doc, "6.2 Figure 8 Product category and review outcomes", 2)
    add_bilingual(
        doc,
        "分析问题：不同产品类别的评分如何变化，评论量是否解释类别结果？观察单位为一条与产品类别匹配的评论。product_reviews通过product_id与products进行many-to-one连接；连接后仍为7,000行、0个孤立键、行数放大倍数1.0。",
        "Analytical question: How do review ratings differ across product categories, and does review volume explain the apparent result? The unit is one review matched to its product category. product_reviews joins products many-to-one on product_id; the validated join retains 7,000 rows, has zero orphan keys and a row-multiplication factor of 1.0.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(figure8_path), width=Inches(6.75))
    add_bilingual(
        doc,
        "结果：平均评分从Audio的3.291到Accessory的4.197，相差0.906星；各类别评论数为666至721，评论量与平均评分仅弱相关，r=0.156。限制：类别均值可能受到产品组合、品牌、价格、履约、评论时点和自选择影响，不能解释因果。",
        "Result: mean rating ranges from 3.291 for Audio to 4.197 for Accessory, a 0.906-star gap. Category counts range from 666 to 721 and volume is weakly correlated with mean rating, r=0.156. Limitation: product mix, brand, price, fulfilment, review timing and reviewer selection may affect category means, so the result is not causal.",
    )

    doc.add_page_break()
    add_heading(doc, "7 Task 6 Findings 9 and 10  发现", 1)
    add_heading(doc, "7.1 Finding 9", 2)
    add_bilingual(
        doc,
        "在一条评论的粒度上，Figure 7包含6,311条英语评论和689条非英语评论。英语评论长度中位数从1星的974字符降至5星的904.5字符，而各评分和语言单元的helpful votes中位数保持在42至49.5。脚本差异、组大小不平衡和评论曝光可能解释该模式；在改变支持政策前，应比较token-normalised length和exposure-adjusted helpfulness。",
        "At the one-review grain, Figure 7 covers 6,311 English and 689 non-English reviews. English median length falls from 974 characters at one star to 904.5 at five stars, while helpful-vote medians across rating/language cells stay within 42–49.5. Script differences, unequal group size and review exposure may explain the pattern; compare token-normalised length and exposure-adjusted helpfulness before changing support policy.",
    )
    add_heading(doc, "7.2 Finding 10", 2)
    add_bilingual(
        doc,
        "Figure 8显示Audio平均评分为3.291（n=707），Accessory为4.197（n=721），差距0.906星。各类别评论数为666至721，并且与平均评分弱相关（r=0.156），因此评论量本身不能解释类别排序。评论者自选择和产品组合仍是合理替代解释；在供应商或产品组合决策前，应进行产品级、时间感知的调整比较。",
        "Figure 8 shows mean rating from 3.291 for Audio (n=707) to 4.197 for Accessory (n=721), a 0.906-star gap. Category counts are balanced at 666–721 and correlate weakly with mean rating (r=0.156), so volume alone does not explain the ordering. Reviewer selection and product mix remain plausible; use product-level, time-aware adjusted comparisons before supplier or assortment decisions.",
    )

    add_heading(doc, "8 Task 6 Machine Learning Questions 4 and 5", 1)
    add_heading(doc, "8.1 MLQ 4 Low review risk", 2)
    add_table(
        doc,
        ["Element", "English response", "中文理解"],
        [
            ["Evidence", "1,166/7,000 reviews (16.7%) are 1–2 stars", "低评分类别不平衡"],
            ["Decision", "Prioritise capacity-limited service recovery after delivery", "交付后优先安排服务补救"],
            ["Type and unit", "Binary classification; one delivered order item", "二元分类；一条已交付order item"],
            ["Target", "Whether a later canonical review has rating <= 2", "后续评论是否为1至2星"],
            ["Predictors", "Product item customer channel carrier service and realised delay available at delivery", "只使用交付时已经知道的特征"],
            ["Validation", "Rolling-origin holdout; PR-AUC capacity-based precision/recall and calibration", "按时间验证，不能只看accuracy"],
            ["Risk", "Exclude rating review text and helpful votes; reviewer selection bias", "排除未来信息并监控选择偏差"],
        ],
        widths=[1.05, 3.65, 2.3],
        font_size=8.4,
    )
    add_heading(doc, "8.2 MLQ 5 Stable review behaviour profiles", 2)
    add_table(
        doc,
        ["Element", "English response", "中文理解"],
        [
            ["Evidence", "0.906-star category gap; median 7 reviews per product, range 1–8", "类别差异明显，但产品评论历史稀疏"],
            ["Decision", "Support product quality and supplier investigation", "帮助质量和供应商调查"],
            ["Type and unit", "Unsupervised clustering; one product at a snapshot", "无监督聚类；一个时间截面的产品"],
            ["Objective", "Stable interpretable groups; no supervised target", "寻找稳定可解释群组，没有target"],
            ["Predictors", "Rating variability length helpfulness language and low-rating shares count price warranty category", "产品属性加历史评论聚合特征"],
            ["Validation", "Earlier/later windows and bootstrap stability; silhouette ARI and Jaccard", "同时检查分离度和时间稳定性"],
            ["Risk", "Sparse history scaling and language proxies can dominate", "样本稀疏、缩放和语言代理风险"],
        ],
        widths=[1.05, 3.65, 2.3],
        font_size=8.35,
    )

    doc.add_page_break()
    add_heading(doc, "9 What I Can Explain Orally  口头讲解提纲", 1)
    for zh, en in [
        ("我负责产品和产品评论域，最终输出分别为1,000行和7,000行。", "I own the product and product-review domain, producing 1,000 Product rows and 7,000 canonical Review rows."),
        ("产品仅来自XML；评论来自JSON和XML，需要处理源内重复和跨源重叠。", "Products are supplied in XML only; reviews occur in both JSON and XML and require within-source and cross-source reconciliation."),
        ("评论中的order reference和SKU必须先从原始文本提取，再清洗正文。", "Embedded order references and SKUs must be extracted from raw text before the narrative is cleaned."),
        ("我的两个Task 3函数使用Unicode脚本信息，而不是ASCII判断。", "My two Task 3 functions use Unicode script information rather than an ASCII test."),
        ("build_latin_analysis生成Latin-only分析文本；contains_non_latin_script保留一个独立布尔信号。", "build_latin_analysis creates a Latin-only analytical representation, while contains_non_latin_script preserves a separate boolean signal."),
        ("我的20项验证全部PASS，但状态来自实际计算，不是硬编码。", "All 20 validations in my scope pass, and the statuses are computed rather than hard-coded."),
        ("Figure 7是不需要join的多变量分组文本分析；Figure 8是经过many-to-one验证的关系分析。", "Figure 7 is a segmented multivariate text analysis without a join; Figure 8 is a validated many-to-one relational analysis."),
        ("MLQ 4是二元分类；MLQ 5是无监督聚类。", "MLQ 4 is binary classification and MLQ 5 is unsupervised clustering."),
    ]:
        add_bullet(doc, zh, en)

    add_heading(doc, "10 Completion and Final Checks  完成情况与待核对事项", 1)
    add_heading(doc, "10.1 Completed  已完成", 2)
    completed = [
        ("Task 2 products标准化与输出", "Task 2 Product standardisation and output"),
        ("Task 2 product_reviews双源标准化、去重、协调与输出", "Dual-source Product Review standardisation, deduplication, reconciliation and output"),
        ("Task 3.4两个多语言函数及边界逻辑", "Both assigned multilingual functions and their boundary logic"),
        ("Task 4产品、评论、文本和最终register范围内20项验证", "Twenty Product, Review, text and final-register validations"),
        ("Figures 7–8及其问题、单位、分母、join、结果与限制", "Figures 7–8 with question, unit, denominator, join, result and limitation"),
        ("Findings 9–10以及MLQ 4–5", "Findings 9–10 and MLQ 4–5"),
        ("个人代码与分析中英讲解材料", "Bilingual explanation of my assigned code and analysis"),
    ]
    for zh, en in completed:
        add_bullet(doc, zh, en)

    add_heading(doc, "10.2 Still requires group confirmation  仍需小组确认", 2)
    checks = [
        ("最终提交目录必须让solution notebook在同级找到text functions、mapping和outputs。", "The final package must allow the solution notebook to find the text module, mapping and outputs through the agreed relative paths."),
        ("从最终提交目录执行Restart Kernel and Run All，并确认六张CSV重新生成。", "Run Restart Kernel and Run All from the final submission directory and confirm that all six CSV files regenerate."),
        ("替换EDA notebook中仍存在的assurance narrative以及limitations/conclusion占位文本。", "Replace the remaining assurance-narrative and limitations/conclusion placeholders in the EDA notebook."),
        ("统一EDA notebook与report中MLQ 4–5的最终措辞。", "Synchronise the final wording of MLQ 4–5 between the EDA notebook and report."),
        ("修正report中MLQ 1的Evidence重复标签，并检查MLQ 5跨页排版。", "Correct the duplicated Evidence label in MLQ 1 and review the page break inside MLQ 5."),
        ("在任务划分表中补全Products/product reviews owner姓名、Task 1、Task 7和Task 8负责人。", "Complete the owner names for Products/Reviews and the currently unassigned Task 1, Task 7 and Task 8 integration work."),
        ("每位成员核对AI-use declaration、最终文件名以及report中的数字与图表。", "Each member must check the AI-use declaration, final filenames, and every reported number and figure."),
    ]
    for zh, en in checks:
        add_bullet(doc, zh, en)

    add_heading(doc, "Source files reviewed  已核对文件", 2)
    add_bullet(doc, "任务划分.docx", "Team responsibility allocation")
    add_bullet(doc, "Group050_solution.ipynb", "Task 2–4 implementation and validation evidence")
    add_bullet(doc, "Group050_text_functions.py", "Final shared Task 3 function module")
    add_bullet(doc, "Group050_EDA.ipynb", "Figures 7–8, Findings 9–10 and MLQ 4–5")
    add_bullet(doc, "Group050_EDA_revised_with_page_numbers.docx", "Latest report alignment reference")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
