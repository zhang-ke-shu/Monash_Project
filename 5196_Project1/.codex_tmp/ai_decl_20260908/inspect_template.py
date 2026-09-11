import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def inches(value):
    return None if value is None else round(value.inches, 4)


path = Path(sys.argv[1])
doc = Document(path)
report = {
    "path": str(path.resolve()),
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "sections": [],
    "paragraphs": [],
    "tables": [],
    "package_parts": [],
}

for section in doc.sections:
    report["sections"].append(
        {
            "page_width_in": inches(section.page_width),
            "page_height_in": inches(section.page_height),
            "left_margin_in": inches(section.left_margin),
            "right_margin_in": inches(section.right_margin),
            "top_margin_in": inches(section.top_margin),
            "bottom_margin_in": inches(section.bottom_margin),
            "header_distance_in": inches(section.header_distance),
            "footer_distance_in": inches(section.footer_distance),
            "start_type": str(section.start_type),
            "different_first_page": section.different_first_page_header_footer,
        }
    )

for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text
    if not text.strip():
        continue
    runs = []
    for run in paragraph.runs:
        if not run.text:
            continue
        rpr = run._element.rPr
        fonts = rpr.rFonts if rpr is not None else None
        runs.append(
            {
                "text": run.text,
                "font": run.font.name,
                "ascii_font": fonts.get(qn("w:ascii")) if fonts is not None else None,
                "size_pt": run.font.size.pt if run.font.size else None,
                "bold": run.bold,
                "italic": run.italic,
                "color": str(run.font.color.rgb) if run.font.color.rgb else None,
            }
        )
    pf = paragraph.paragraph_format
    report["paragraphs"].append(
        {
            "index": index,
            "text": text,
            "style": paragraph.style.name,
            "alignment": str(paragraph.alignment),
            "space_before_pt": pf.space_before.pt if pf.space_before else None,
            "space_after_pt": pf.space_after.pt if pf.space_after else None,
            "line_spacing": pf.line_spacing,
            "left_indent_in": inches(pf.left_indent),
            "first_line_indent_in": inches(pf.first_line_indent),
            "runs": runs,
        }
    )

for table_index, table in enumerate(doc.tables):
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for col_index, cell in enumerate(row.cells):
            cells.append(
                {
                    "row": row_index,
                    "column": col_index,
                    "width_in": inches(cell.width),
                    "text": "\n".join(p.text for p in cell.paragraphs),
                    "paragraph_styles": [p.style.name for p in cell.paragraphs],
                }
            )
        rows.append(cells)
    report["tables"].append(
        {
            "index": table_index,
            "style": table.style.name if table.style else None,
            "rows": rows,
        }
    )

with zipfile.ZipFile(path) as package:
    for info in package.infolist():
        data = package.read(info.filename)
        report["package_parts"].append(
            {
                "path": info.filename,
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
        )

print(json.dumps(report, ensure_ascii=False, indent=2))
