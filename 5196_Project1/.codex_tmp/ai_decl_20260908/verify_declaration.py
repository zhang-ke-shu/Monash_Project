import hashlib
import sys
import zipfile
from pathlib import Path

from docx import Document


reference = Path(sys.argv[1])
final = Path(sys.argv[2])

expected_reference = "2577cfe8f0cc7d997dbaaed81ea1e848b5ce0b98fe6d5a3f927068d3ede71286"
actual_reference = hashlib.sha256(reference.read_bytes()).hexdigest()
assert actual_reference == expected_reference, actual_reference

with zipfile.ZipFile(reference) as source, zipfile.ZipFile(final) as output:
    assert source.namelist() == output.namelist(), "Package-part list or order changed"
    changed = []
    for name in source.namelist():
        if source.read(name) != output.read(name):
            changed.append(name)
    assert changed == ["word/document.xml"], changed

doc = Document(final)
body = "\n".join(p.text for p in doc.paragraphs)
assert "GroupNNN" not in body
assert "Group Number: Group050" in body
assert doc.tables[0].cell(1, 1).text.strip() == "Keshu Zhang"
assert doc.tables[0].cell(1, 2).text.strip() == "36436763"
assert "OpenAI Codex (GPT-5)" in doc.tables[1].cell(1, 1).text
assert "AI01_Keshu.pdf" in doc.tables[1].cell(1, 1).text
assert "Figures 7-8" in doc.tables[1].cell(1, 2).text
assert "20 passing validations" in doc.tables[1].cell(1, 3).text
assert all(not cell.text.strip() for cell in doc.tables[2].rows[1].cells)
assert all(not cell.text.strip() for cell in doc.tables[2].rows[2].cells)

print("Reference hash unchanged")
print("Only word/document.xml changed")
print("Required declaration fields verified")
