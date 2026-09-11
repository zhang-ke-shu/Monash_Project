from pathlib import Path
import sys

from docx import Document


def show(path: str) -> None:
    print(f"\n===== {path} =====")
    doc = Document(path)
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"P{index}: {text}")
    for table_index, table in enumerate(doc.tables):
        print(f"TABLE {table_index}:")
        for row_index, row in enumerate(table.rows):
            values = [
                " | ".join(p.text for p in cell.paragraphs).strip()
                for cell in row.cells
            ]
            print(f"R{row_index}: " + " || ".join(values))


for source in sys.argv[1:]:
    show(str(Path(source)))
